#!/usr/bin/env python3
"""
Simple Markdown-to-Word exporter for COMPLETE_SYSTEM_REPORT.md.
Usage: python tools/export_report_to_word.py

Requires: python-docx
Install: pip install python-docx
"""
from docx import Document
from docx.shared import Pt
from pathlib import Path

SRC = Path(__file__).resolve().parents[1] / 'COMPLETE_SYSTEM_REPORT.md'
OUT = Path(__file__).resolve().parents[1] / 'COMPLETE_SYSTEM_REPORT.docx'

def add_code_paragraph(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    font = run.font
    font.name = 'Consolas'
    font.size = Pt(9)


def md_to_docx(src: Path, out: Path):
    doc = Document()
    if not src.exists():
        raise FileNotFoundError(f"Source not found: {src}")

    with src.open('r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code = False
    for raw in lines:
        line = raw.rstrip('\n')
        if line.strip().startswith('```'):
            in_code = not in_code
            continue
        if in_code:
            add_code_paragraph(doc, line)
            continue
        if not line.strip():
            # blank line -> add empty paragraph to separate
            doc.add_paragraph('')
            continue
        if line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=1)
            continue
        if line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
            continue
        if line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
            continue
        if line.lstrip().startswith('- '):
            # bullet
            text = line.lstrip()[2:].strip()
            doc.add_paragraph(text, style='List Bullet')
            continue
        # fallback: plain paragraph
        doc.add_paragraph(line)

    doc.save(out)
    print(f'Wrote: {out}')


if __name__ == '__main__':
    md_to_docx(SRC, OUT)
